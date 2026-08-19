from __future__ import annotations

import os
import sys
from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor, Twips


ROOT = Path(__file__).resolve().parent
OUT_DIR = ROOT / "outputs" / "reporte_vtex_vs_tiendanube_2026-08-18"
OUT_DIR.mkdir(parents=True, exist_ok=True)
OUTPUT = OUT_DIR / "Reporte_B2C_VTEX_vs_Tiendanube_Evolucion_vs_Shopify_Saracho_Neumaticos.docx"

DOC_SKILL = Path(
    r"C:\Users\Joaco\.codex\plugins\cache\openai-primary-runtime\documents\26.813.12317\skills\documents"
)
sys.path.insert(0, str(DOC_SKILL / "scripts"))
from table_geometry import apply_table_geometry  # noqa: E402


NAVY = "17365D"
BLUE = "2E74B5"
TEAL = "178A8A"
GREEN = "2E7D32"
AMBER = "B36B00"
RED = "B3261E"
INK = "1F2937"
MID = "5B6573"
LIGHT = "F2F4F7"
PALE_BLUE = "EAF2F8"
PALE_GREEN = "EAF5EC"
PALE_AMBER = "FFF4DE"
PALE_RED = "FDECEA"
WHITE = "FFFFFF"
RULE = "D6DEE8"


def rgb(hex_value: str) -> RGBColor:
    return RGBColor.from_string(hex_value)


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_border(cell, **kwargs) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_borders = tc_pr.first_child_found_in("w:tcBorders")
    if tc_borders is None:
        tc_borders = OxmlElement("w:tcBorders")
        tc_pr.append(tc_borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        edge_data = kwargs.get(edge)
        if not edge_data:
            continue
        tag = "w:{}".format(edge)
        element = tc_borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            tc_borders.append(element)
        for key in ("val", "sz", "space", "color"):
            if key in edge_data:
                element.set(qn("w:{}".format(key)), str(edge_data[key]))


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_keep_with_next(paragraph, value: bool = True) -> None:
    paragraph.paragraph_format.keep_with_next = value


def set_keep_together(paragraph, value: bool = True) -> None:
    paragraph.paragraph_format.keep_together = value


def set_font(run, size=None, bold=None, italic=None, color=INK, name="Calibri") -> None:
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    if color:
        run.font.color.rgb = rgb(color)


def add_hyperlink(paragraph, text: str, url: str, color=BLUE, underline=True):
    part = paragraph.part
    relationship_id = part.relate_to(
        url,
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), relationship_id)
    new_run = OxmlElement("w:r")
    r_pr = OxmlElement("w:rPr")
    c = OxmlElement("w:color")
    c.set(qn("w:val"), color)
    r_pr.append(c)
    if underline:
        u = OxmlElement("w:u")
        u.set(qn("w:val"), "single")
        r_pr.append(u)
    r_fonts = OxmlElement("w:rFonts")
    r_fonts.set(qn("w:ascii"), "Calibri")
    r_fonts.set(qn("w:hAnsi"), "Calibri")
    r_pr.append(r_fonts)
    new_run.append(r_pr)
    text_element = OxmlElement("w:t")
    text_element.text = text
    new_run.append(text_element)
    hyperlink.append(new_run)
    paragraph._p.append(hyperlink)
    return hyperlink


def add_field(paragraph, instruction: str, fallback: str = "1") -> None:
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = instruction
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = fallback
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    for node in (begin, instr, separate, text, end):
        run._r.append(node)
    set_font(run, size=8.5, color=MID)


def paragraph_rule(paragraph, color=RULE, size="10", space="6") -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    p_bdr = p_pr.find(qn("w:pBdr"))
    if p_bdr is None:
        p_bdr = OxmlElement("w:pBdr")
        p_pr.append(p_bdr)
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), size)
    bottom.set(qn("w:space"), space)
    bottom.set(qn("w:color"), color)
    p_bdr.append(bottom)


def shade_paragraph(paragraph, fill: str) -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    p_pr.append(shd)


def set_doc_defaults(doc: Document) -> None:
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.78)
    section.bottom_margin = Inches(0.74)
    section.left_margin = Inches(1.0)
    section.right_margin = Inches(1.0)
    section.header_distance = Inches(0.32)
    section.footer_distance = Inches(0.32)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = rgb(INK)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.10

    for name, size, color, before, after in (
        ("Title", 25, NAVY, 0, 5),
        ("Subtitle", 13.5, MID, 0, 12),
        ("Heading 1", 16, BLUE, 16, 8),
        ("Heading 2", 13, BLUE, 12, 6),
        ("Heading 3", 11.5, NAVY, 8, 4),
    ):
        style = styles[name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
        style.font.size = Pt(size)
        style.font.color.rgb = rgb(color)
        style.font.bold = True
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    for name in ("List Bullet", "List Number"):
        style = styles[name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
        style.font.size = Pt(10.5)
        style.paragraph_format.left_indent = Inches(0.5)
        style.paragraph_format.first_line_indent = Inches(-0.25)
        style.paragraph_format.space_after = Pt(5)
        style.paragraph_format.line_spacing = 1.10


def set_running_header_footer(section) -> None:
    header = section.header
    p = header.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_after = Pt(2)
    run = p.add_run("SARACHO NEUMÁTICOS  |  INFORME DE DECISIÓN")
    set_font(run, size=8.3, bold=True, color=MID)
    paragraph_rule(p, color=RULE, size="6", space="3")

    footer = section.footer
    p = footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p.paragraph_format.space_before = Pt(2)
    run = p.add_run("Agosto 2026   •   ")
    set_font(run, size=8.3, color=MID)
    add_field(p, "PAGE", "1")


def add_text(doc, text: str, *, bold_prefix: str | None = None, italic=False, after=6, align=None):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(after)
    if align is not None:
        p.alignment = align
    if bold_prefix and text.startswith(bold_prefix):
        r = p.add_run(bold_prefix)
        set_font(r, bold=True)
        r = p.add_run(text[len(bold_prefix):])
        set_font(r, italic=italic)
    else:
        r = p.add_run(text)
        set_font(r, italic=italic)
    return p


def add_bullet(doc, text: str, *, level=0, bold_prefix: str | None = None):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.left_indent = Inches(0.5 + level * 0.25)
    p.paragraph_format.first_line_indent = Inches(-0.25)
    if bold_prefix and text.startswith(bold_prefix):
        r = p.add_run(bold_prefix)
        set_font(r, bold=True)
        r = p.add_run(text[len(bold_prefix):])
        set_font(r)
    else:
        r = p.add_run(text)
        set_font(r)
    return p


def add_number(doc, text: str, *, bold_prefix: str | None = None):
    p = doc.add_paragraph(style="List Number")
    if bold_prefix and text.startswith(bold_prefix):
        r = p.add_run(bold_prefix)
        set_font(r, bold=True)
        r = p.add_run(text[len(bold_prefix):])
        set_font(r)
    else:
        r = p.add_run(text)
        set_font(r)
    return p


def add_callout(doc, label: str, text: str, fill: str, accent: str) -> None:
    table = doc.add_table(rows=1, cols=1)
    set_repeat_table_header(table.rows[0])
    cell = table.cell(0, 0)
    set_cell_shading(cell, fill)
    set_cell_border(
        cell,
        left={"val": "single", "sz": "26", "color": accent, "space": "0"},
        top={"val": "nil"},
        bottom={"val": "nil"},
        right={"val": "nil"},
    )
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(2)
    r = p.add_run(label.upper())
    set_font(r, size=9.2, bold=True, color=accent)
    p = cell.add_paragraph()
    p.paragraph_format.space_after = Pt(1)
    p.paragraph_format.line_spacing = 1.08
    r = p.add_run(text)
    set_font(r, size=10.8, bold=True, color=INK)
    apply_table_geometry(
        table,
        [9360],
        table_width_dxa=9360,
        indent_dxa=180,
        cell_margins_dxa={"top": 130, "bottom": 130, "start": 180, "end": 180},
    )
    doc.add_paragraph().paragraph_format.space_after = Pt(1)


def add_table(doc, headers, rows, widths, *, font_size=8.7, first_col_bold=False, status_colors=None):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    hdr = table.rows[0]
    set_repeat_table_header(hdr)
    for i, value in enumerate(headers):
        cell = hdr.cells[i]
        set_cell_shading(cell, NAVY)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT if i else WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.space_after = Pt(0)
        r = p.add_run(str(value))
        set_font(r, size=8.5, bold=True, color=WHITE)
    for row_idx, row in enumerate(rows):
        cells = table.add_row().cells
        for i, value in enumerate(row):
            cell = cells[i]
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if row_idx % 2 == 1:
                set_cell_shading(cell, "F8FAFC")
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.03
            r = p.add_run(str(value))
            color = INK
            bold = first_col_bold and i == 0
            if status_colors and i in status_colors:
                for key, status_color in status_colors[i].items():
                    if key.lower() in str(value).lower():
                        color = status_color
                        bold = True
                        break
            set_font(r, size=font_size, bold=bold, color=color)
    apply_table_geometry(
        table,
        widths,
        table_width_dxa=9360,
        indent_dxa=120,
        cell_margins_dxa={"top": 90, "bottom": 90, "start": 120, "end": 120},
    )
    for row in table.rows:
        row.height = None
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return table


def add_small_source(doc, text: str, url: str | None = None):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(4)
    r = p.add_run(text)
    set_font(r, size=8.1, italic=True, color=MID)
    if url:
        add_hyperlink(p, " Abrir fuente", url)
    return p


def title_page(doc: Document) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(18)
    p.paragraph_format.space_after = Pt(4)
    r = p.add_run("INFORME DE DECISIÓN")
    set_font(r, size=10.2, bold=True, color=TEAL)

    p = doc.add_paragraph(style="Title")
    p.paragraph_format.space_after = Pt(4)
    r = p.add_run("VTEX vs. Tiendanube vs. Shopify")
    set_font(r, size=27, bold=True, color=NAVY)

    p = doc.add_paragraph(style="Subtitle")
    p.paragraph_format.space_after = Pt(12)
    r = p.add_run("Evaluación B2C de plataformas, integradores, promociones, costos y capacidad operativa")
    set_font(r, size=14, color=MID)

    meta_rows = [
        ("Empresa", "Saracho Neumáticos"),
        ("Fecha", "18 de agosto de 2026"),
        ("Objetivo", "Definir una ruta B2C de plataforma e integración con menor costo y riesgo operativo"),
        ("Alcance", "Experiencia B2C, catálogo, promociones, cuotas, pagos, stock, pedidos, diseño y soporte"),
        ("Estado", "Recomendación preliminar sujeta a costos reales y prueba técnica"),
    ]
    table = doc.add_table(rows=0, cols=2)
    table.style = "Table Grid"
    for label, value in meta_rows:
        cells = table.add_row().cells
        set_cell_shading(cells[0], PALE_BLUE)
        for idx, text in enumerate((label, value)):
            p = cells[idx].paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            r = p.add_run(text)
            set_font(r, size=9.3, bold=(idx == 0), color=NAVY if idx == 0 else INK)
            cells[idx].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    set_repeat_table_header(table.rows[0])
    apply_table_geometry(table, [1900, 7460], table_width_dxa=9360, indent_dxa=120)
    doc.add_paragraph().paragraph_format.space_after = Pt(3)

    add_callout(
        doc,
        "Conclusión ejecutiva",
        "Saracho es un negocio B2C. La prioridad es bajar el costo fijo porque el nivel de ventas está bajo, sin resignar conversión, imagen ni capacidad comercial. La dotación web es reducida y existen promociones complejas —por ejemplo, un 4x3 válido únicamente en 6 cuotas— que la implementación actual de VTEX no permite limitar así. Tiendanube debe evaluarse específicamente bajo el plan Evolución; Shopify es la tercera alternativa. Astroselling–VTEX queda como hipótesis.",
        PALE_AMBER,
        AMBER,
    )

    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(7)
    p.paragraph_format.space_after = Pt(3)
    r = p.add_run("DECISIÓN PROPUESTA")
    set_font(r, size=9.2, bold=True, color=TEAL)
    for text in (
        "Mantener el esquema actual durante el análisis y no hacer un reemplazo directo de Producteca.",
        "Construir demos comparables de Tiendanube Evolución y Shopify —Home, listado, ficha, cuotas, carrito y móvil— para medir calidad visual y experiencia contra VTEX.",
        "Cotizar Tiendanube Evolución con condiciones B2C reales: fee, transacciones, Pago Nube, soporte, migración, desarrollo custom y aplicaciones.",
        "Probar la matriz promocional real, incluyendo un 4x3 habilitado solo para 6 cuotas, combinaciones permitidas, exclusiones y comunicación al cliente.",
        "Medir si una persona puede operar catálogo, precios, promociones y controles sin depender continuamente de desarrollo o múltiples proveedores.",
        "Comparar tres alternativas principales: continuidad optimizada en VTEX + Producteca, Tiendanube Evolución + Astroselling y Shopify + Astroselling, usando TCO a 24 meses, impacto inmediato en caja y pruebas técnica/visual.",
        "Mantener la posible integración Astroselling–VTEX como caso hipotético y no incluir sus ahorros en el business case hasta que exista una propuesta verificable.",
    ):
        add_bullet(doc, text)

    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(7)
    p.paragraph_format.space_after = Pt(2)
    r = p.add_run("Lectura rápida")
    set_font(r, size=9.2, bold=True, color=NAVY)
    add_text(
        doc,
        "Por ser B2C, pesan especialmente la conversión móvil, la búsqueda por medida, cuotas, promociones, checkout, pagos, envíos y autogestión. Capacidades B2B, marketplace propio u OMS enterprise solo agregan valor si existe un caso concreto. Tiendanube Evolución aporta acompañamiento y personalización, pero al cotizarse a medida no debe presumirse más barata que VTEX o Shopify.",
        after=0,
    )


def section_context(doc: Document) -> None:
    doc.add_page_break()
    doc.add_heading("1. Contexto y arquitectura actual", level=1)
    add_text(
        doc,
        "Saracho Neumáticos opera un negocio B2C sobre VTEX. Producteca sincroniza el catálogo con VTEX y, por lo tanto, interviene en datos críticos como productos, precios, stock y pedidos. Pirce se incorpora como integrador/partner operativo actual; su alcance, contrato, fee y responsabilidades deben validarse antes de decidir.",
    )
    add_text(
        doc,
        "La documentación interna también evidencia desarrollo específico sobre VTEX IO para información de cuotas y promociones. Ese activo no migra automáticamente a Tiendanube Evolución ni Shopify: deberá reemplazarse por funcionalidad nativa, una aplicación equivalente o un nuevo desarrollo.",
    )
    add_text(
        doc,
        "Como antecedente comercial relevante, Saracho Neumáticos ya mantuvo conversaciones con Astroselling para analizar una integración con VTEX y el proveedor manifestó disposición a desarrollarla. El informe trata esa disposición como una oportunidad concreta, aunque pendiente de propuesta formal, definición técnica, costos, SLA y prueba productiva.",
    )
    add_text(
        doc,
        "La operación enfrenta además una restricción comercial concreta: existen promociones que combinan beneficio y financiación, como un 4x3 aplicable únicamente si el cliente paga en 6 cuotas. Según la experiencia actual de Saracho, VTEX no permite limitar esa promoción por cantidad de cuotas de la forma requerida, lo que obliga a simplificar la oferta, asumir excepciones o sumar desarrollo.",
    )
    add_text(
        doc,
        "El equipo dedicado a la página web es reducido. Esto aumenta el costo relativo de cada tarea manual, capacitación e incidente, y vuelve riesgosa cualquier arquitectura que dependa de varios paneles, controles diarios o escalamiento entre plataforma, integrador y desarrollador.",
    )
    add_callout(
        doc,
        "Prioridad del negocio",
        "La baja de ventas vuelve urgente reducir el punto de equilibrio mensual. Con una dotación reducida, la decisión debe privilegiar ahorro recurrente, autogestión y menos tareas manuales, pero proteger conversión, promociones, posicionamiento de marca y continuidad operativa.",
        PALE_GREEN,
        GREEN,
    )
    add_callout(
        doc,
        "Lente B2C",
        "La plataforma ganadora debe facilitar que un consumidor encuentre el neumático correcto, entienda precio y financiación, aplique la promoción válida y complete la compra desde el celular. Las capacidades enterprise que no mejoren ese recorrido ni reduzcan trabajo operativo deben considerarse costo excedente.",
        PALE_BLUE,
        BLUE,
    )

    doc.add_heading("Mapa simplificado", level=2)
    rows = [
        ("VTEX", "Tienda, catálogo operativo, promociones, checkout, pedidos y extensiones VTEX IO", "Plataforma central"),
        ("Producteca", "Sincronización multicanal de catálogo, stock, precios y pedidos", "Capa crítica de integración"),
        ("Pirce", "Implementación, soporte o integración según contrato vigente", "Responsabilidad a confirmar"),
        ("Mercado Libre y otros", "Canales externos que consumen catálogo/stock/precio y generan ventas", "Canales comerciales"),
    ]
    add_table(doc, ["Componente", "Función considerada", "Rol"], rows, [1900, 5560, 1900], font_size=8.8, first_col_bold=True)

    add_callout(
        doc,
        "Principio de arquitectura",
        "Cada salto adicional entre sistemas agrega puntos de falla, conciliaciones y dependencia de terceros. La decisión no es solo VTEX vs. Tiendanube vs. Shopify: es qué combinación completa tiene un dueño claro para cada dato y cada incidente.",
        PALE_BLUE,
        BLUE,
    )

    doc.add_heading("Supuestos y límites", level=2)
    for text in (
        "El análisis es ejecutivo y no reemplaza una cotización vinculante ni un relevamiento funcional detallado.",
        "No se encontraron referencias públicas inequívocas sobre Pirce; se evita atribuirle funciones no confirmadas.",
        "Los precios publicados cambian —especialmente en Argentina— y deben revalidarse al contratar.",
        "No se dispone aún del GMV, cantidad de pedidos, SKU activos, centros de stock, fee de cada proveedor ni horas mensuales de soporte.",
    ):
        add_bullet(doc, text)


def section_platform_comparison(doc: Document) -> None:
    doc.add_heading("2. Comparación de plataformas", level=1)
    add_text(
        doc,
        "La comparación parte de un negocio B2C. VTEX está orientada a comercio unificado y extensible; Tiendanube Evolución ofrece una propuesta a medida con personalización, migración y acompañamiento; Shopify combina una operación SaaS accesible con un ecosistema global de temas, aplicaciones y desarrollo.",
    )

    rows = [
        ("Implementación", "Mayor esfuerzo y especialistas", "Más rápida y local", "Rápida con tema; custom exige partner"),
        ("Costo", "Contrato + Producteca + Pirce + mantenimiento", "Evolución se cotiza a medida + apps/pagos", "USD + apps + fee por pagos externos"),
        ("Ajuste B2C", "Muy capaz; puede estar sobredimensionada", "Foco local, conversión, pagos y soporte dedicado", "Gran UX, temas y ecosistema global"),
        ("Complejidad", "Alta escala, OMS y reglas", "Operación estándar/local", "Escala SaaS y ecosistema global"),
        ("Personalización", "VTEX IO, APIs y headless", "HTML/CSS/JS según plan", "Temas, Liquid, apps, APIs y headless"),
        ("Equipo reducido", "Más coordinación y conocimiento", "Administración simple; validar custom", "Panel accesible y soporte 24/7; apps agregan terceros"),
        ("Marketplaces", "Producteca centraliza hoy", "Astroselling documentado", "Astroselling documentado"),
        ("4x3 + 6 cuotas", "No resuelto en la operación actual", "Debe demostrarse; posible app/custom", "Buy X Get Y nativo, pero condicionar por cuotas requiere validación/custom"),
        ("Pagos Argentina", "Ecosistema actual integrado", "Ventaja de soluciones locales", "Shopify Payments no figura para Argentina; proveedor externo"),
        ("Omnicanal/OMS", "Ventaja en operaciones complejas", "Más acotado", "Apps/POS; no equivale al OMS VTEX"),
        ("Riesgo", "Costo y dependencia técnica", "Límites y dependencia de apps", "Dolarización, pagos externos y costo de apps"),
    ]
    add_table(doc, ["Criterio", "VTEX", "Tiendanube Evolución", "Shopify"], rows, [1500, 2620, 2620, 2620], font_size=7.35, first_col_bold=True)

    doc.add_heading("Ventajas principales de VTEX", level=2)
    for text in (
        "Escalabilidad y control: mejor preparada para reglas comerciales, inventario, OMS, marketplaces y operaciones complejas.",
        "Extensibilidad: permite construir diferenciadores mediante VTEX IO, APIs y opciones composables/headless.",
        "Continuidad: evita una migración de catálogo, SEO, promociones, clientes, analytics y procesos ya estabilizados.",
        "Aprovechamiento de activos existentes: conserva la app personalizada de cuotas/promociones y el conocimiento acumulado.",
    ):
        add_bullet(doc, text)

    doc.add_heading("Desventajas principales de VTEX", level=2)
    for text in (
        "Costo total y complejidad más altos: plataforma, Producteca, Pirce, aplicaciones y horas de desarrollo deben analizarse como un conjunto.",
        "Mayor dependencia de especialistas: cambios y diagnósticos pueden involucrar a más de un proveedor.",
        "Limitación promocional observada: la configuración actual no permite condicionar una promoción como 4x3 exclusivamente al pago en 6 cuotas.",
        "Carga operativa para un equipo reducido: administrar reglas, integraciones e incidentes puede concentrarse en pocas personas.",
        "Riesgo de sobredimensionamiento: si la operación es esencialmente una tienda B2C estándar más Mercado Libre, parte de la capacidad puede no generar retorno.",
        "Gobernanza exigente: se necesita definir qué sistema manda sobre catálogo, stock, precio y pedidos.",
    ):
        add_bullet(doc, text)

    doc.add_heading("Ventajas principales de Tiendanube Evolución", level=2)
    for text in (
        "Menor barrera operativa: puesta en marcha, administración y capacitación más simples.",
        "Propuesta integral: el plan puede reunir plataforma, acompañamiento, migración y condiciones comerciales específicas para la marca.",
        "Acompañamiento: Evolución publica especialista dedicado, soporte prioritario y migración asistida, valiosos para una dotación reducida.",
        "Ecosistema local B2C: pagos, cuotas, envíos y aplicaciones pensadas para negocios argentinos.",
        "Diseño y desarrollo: Evolución publica acceso al front-end, APIs y personalización avanzada.",
        "Compatibilidad documentada con Astroselling: sincronización de precio y stock, publicación y centralización de ventas multicanal.",
    ):
        add_bullet(doc, text)

    doc.add_heading("Desventajas principales de Tiendanube Evolución", level=2)
    for text in (
        "Menor profundidad para casos complejos: hay que validar promociones, múltiples inventarios, fulfillment, permisos y flujos especiales.",
        "Precio no público: Evolución se cotiza a medida; con ventas bajas, debe demostrar un ahorro real frente al TCO actual.",
        "La combinación 4x3 + 6 cuotas no debe darse por resuelta: puede requerir una aplicación, una personalización o una adaptación de la mecánica comercial.",
        "Dependencia de aplicaciones: el costo y la confiabilidad final no son solo los del plan base.",
        "Comisiones: con medios de pago externos, los planes publicados contemplan costo por transacción.",
        "Costo de migración y posible pérdida funcional: SEO, customizaciones, históricos y procesos deben reconstruirse o resignarse.",
    ):
        add_bullet(doc, text)

    doc.add_heading("Ventajas principales de Shopify", level=2)
    for text in (
        "Diseño y ecosistema: amplia oferta de temas, aplicaciones, partners y opciones de desarrollo custom.",
        "Operación accesible: panel integrado, productos ilimitados y soporte 24/7 en español publicado por Shopify.",
        "Promociones base: incluye descuentos Compra X y obtén Y, útiles como punto de partida para un 4x3.",
        "Compatibilidad con Astroselling: Shopify figura entre las integraciones oficiales publicadas por Astroselling.",
    ):
        add_bullet(doc, text)

    doc.add_heading("Desventajas principales de Shopify", level=2)
    for text in (
        "Costos dolarizados: plan, temas, aplicaciones y desarrollo quedan expuestos al tipo de cambio.",
        "Pagos en Argentina: Argentina no figura en la lista oficial de Shopify Payments, por lo que debe utilizarse un proveedor externo y contemplar el cargo adicional de Shopify según el plan.",
        "Promoción 4x3 + 6 cuotas: el Compra X y obtén Y nativo no demuestra por sí solo la condición sobre cantidad de cuotas; puede requerir app o desarrollo con Shopify Functions.",
        "Dependencia del ecosistema: una tienda simple puede terminar sumando varias aplicaciones, costos y responsables.",
    ):
        add_bullet(doc, text)

    doc.add_heading("Caso crítico: promociones condicionadas por cuotas", level=2)
    add_text(
        doc,
        "La promoción debe evaluarse como un flujo completo y no solo como un descuento. El requisito de referencia es: el cliente compra cuatro neumáticos, recibe el beneficio 4x3 únicamente si elige 6 cuotas, y no debe obtenerlo con otra financiación. Ninguna alternativa se considera apta hasta demostrar este comportamiento de punta a punta.",
    )
    promo_rows = [
        ("Elegibilidad", "Cuatro unidades/SKU o familia válidos; definir si admite mezcla de medidas"),
        ("Financiación", "Beneficio activo solo con 6 cuotas y medio/tarjeta habilitados"),
        ("Exclusiones", "No combinar con descuentos incompatibles, cupones o listas especiales"),
        ("Comunicación", "Condición visible y consistente en PLP, PDP, carrito y checkout"),
        ("Operación", "Crear, pausar y auditar la promo sin intervención técnica recurrente"),
        ("Control", "Pedido final conserva beneficio, cuotas y trazabilidad correctos"),
    ]
    add_table(doc, ["Dimensión", "Condición que debe demostrarse"], promo_rows, [2000, 7360], font_size=8.55, first_col_bold=True)

    doc.add_heading("Operación con una dotación reducida", level=2)
    for text in (
        "Priorizar un único panel o un flujo claramente documentado para catálogo, precio, stock, promociones y pedidos.",
        "Evitar controles manuales diarios y conciliaciones entre varias fuentes de verdad.",
        "Definir un solo dueño del incidente aunque participen plataforma, Astroselling/Producteca y desarrollador.",
        "Medir horas semanales, cantidad de pasos por cambio, errores y dependencia de soporte externo.",
        "No elegir una solución más barata si traslada el ahorro a trabajo manual que el equipo no puede absorber.",
    ):
        add_bullet(doc, text)

    doc.add_heading("Demostración visual custom en Tiendanube Evolución y Shopify", level=2)
    add_text(
        doc,
        "La comparación no debe hacerse entre una VTEX trabajada durante años y plantillas básicas. Para que la decisión sea justa, se propone construir demos comparables en Tiendanube Evolución y Shopify con la identidad real de Saracho y medirlas contra las pantallas actuales y el recorrido B2C móvil.",
    )
    demo_rows = [
        ("Home", "Marca, propuesta comercial, financiación, categorías y acceso rápido por medida"),
        ("Listado / PLP", "Filtros claros, tarjetas de producto, precio, cuotas, stock y navegación móvil"),
        ("Ficha / PDP", "Medida, atributos, beneficios, envío, cuotas, confianza y llamado a compra"),
        ("Carrito y checkout", "Continuidad visual, claridad de costos, pagos, envío y ausencia de fricción"),
        ("Móvil", "Jerarquía, velocidad percibida, botones, filtros y lectura sin pérdida de identidad"),
        ("SEO y analítica", "Metadatos, URLs, datos estructurados, eventos y medición comparables"),
    ]
    add_table(doc, ["Pantalla/tema", "Qué debe demostrar la demo custom"], demo_rows, [2100, 7260], font_size=8.55, first_col_bold=True)
    add_small_source(
        doc,
        "Tiendanube Evolución publica diseño 100% personalizable, acceso al front-end, APIs y soluciones a medida; el alcance contractual debe quedar incluido en la propuesta.",
        "https://www.tiendanube.com/evolucion/plataforma",
    )
    add_small_source(
        doc,
        "Shopify publica temas, tienda online completa y opciones de personalización; las funciones y costos dependen del plan, las apps y el alcance custom.",
        "https://www.shopify.com/ar/precios",
    )


def section_ecosystems(doc: Document) -> None:
    doc.add_heading("3. Evaluación de los ecosistemas posibles", level=1)
    rows = [
        (
            "A. VTEX + Producteca + Pirce",
            "Integración VTEX documentada; continuidad de customizaciones y operación",
            "Costo, complejidad y coordinación de tres proveedores",
            "Viable / base actual",
        ),
        (
            "B. Tiendanube Evolución + Astroselling",
            "Propuesta local a medida, soporte dedicado, migración y conexión documentada",
            "Precio a cotizar; demostrar ahorro, 4x3 + 6 cuotas y diseño",
            "Alternativa a evaluar",
        ),
        (
            "C. Shopify + Astroselling",
            "Buen diseño, ecosistema amplio y conexión Astroselling publicada",
            "USD, pagos externos en Argentina, apps y promo por cuotas",
            "Alternativa a evaluar",
        ),
        (
            "D. Evolución/Shopify + Producteca",
            "Permite desacoplar la migración de plataforma del cambio de integrador",
            "Mantiene temporalmente el costo de Producteca y requiere doble transición",
            "Útil como contingencia",
        ),
    ]
    add_table(
        doc,
        ["Escenario", "Fortaleza", "Riesgo principal", "Juicio"],
        rows,
        [2350, 2850, 2960, 1200],
        font_size=7.85,
        first_col_bold=True,
        status_colors={3: {"Viable": GREEN, "evaluar": AMBER, "contingencia": BLUE}},
    )

    add_callout(
        doc,
        "Punto de decisión",
        "Las alternativas principales son Tiendanube Evolución + Astroselling y Shopify + Astroselling frente a VTEX + Producteca. Por ser B2C, deben demostrar ahorro, promoción 4x3 + 6 cuotas, recorrido móvil, pagos/envíos, operación con poco personal y calidad visual custom.",
        PALE_BLUE,
        BLUE,
    )

    add_callout(
        doc,
        "Caso hipotético: VTEX + Astroselling",
        "Astroselling manifestó disposición a desarrollar una integración con VTEX, pero todavía no existe un conector público probado, propuesta formal, costo, SLA ni alcance confirmado. Por eso este escenario no integra el business case base; solo se reabre si el proveedor presenta una oferta verificable y una prueba satisfactoria.",
        LIGHT,
        MID,
    )

    doc.add_heading("Producteca vs. Astroselling", level=2)
    rows = [
        ("VTEX", "Integración específica publicada", "Disposición comunicada a Saracho para desarrollarla; no figura como conector estándar público", "Potencial a validar"),
        ("Tiendanube Evolución", "Integración disponible", "Integración documentada por Tiendanube y Astroselling", "Ambos viables"),
        ("Shopify", "Integración disponible", "Shopify figura entre las integraciones publicadas por Astroselling", "Ambos viables"),
        ("Catálogo/stock/precio", "Sincronización y publicación multicanal", "Sincronización, reglas de stock/precio, publicación y enlaces", "Validar equivalencia"),
        ("Pedidos/facturación", "Órdenes, envíos y facturas dentro de una plataforma de integración", "Centraliza órdenes; factura si existe ERP/facturador conectado", "Validar flujo real"),
        ("Monitoreo/SLA", "Producteca publica centro de monitoreo, ISO 27001 y SLA según plan", "Revisar SLA, soporte, auditoría y recuperación contractual", "Ventaja a comprobar"),
        ("Precio público", "No fijo; términos contemplan implementación, fee, usuarios/canales y desarrollo", "Planes publicados en USD con límites por canales, SKU y órdenes", "Ventaja Astroselling"),
    ]
    add_table(doc, ["Tema", "Producteca", "Astroselling", "Lectura"], rows, [1650, 3000, 3210, 1500], font_size=7.8, first_col_bold=True)

    add_small_source(
        doc,
        "Fuentes oficiales: Producteca publica su integración con VTEX; Astroselling publica conectores para Tiendanube y Shopify; Tiendanube documenta la instalación de Astroselling. La disposición a integrar VTEX surge únicamente de conversaciones mantenidas por Saracho Neumáticos y se trata como hipótesis.",
    )


def section_costs(doc: Document) -> None:
    doc.add_page_break()
    doc.add_heading("4. Costos: comparar TCO, no solo licencias", level=1)
    add_text(
        doc,
        "La comparación correcta es a 24 meses y debe incluir costos visibles y ocultos. Un plan más económico puede resultar más caro si exige desarrollo, conciliaciones manuales, pérdida de ventas por stock incorrecto o recreación de funcionalidades.",
    )
    add_callout(
        doc,
        "Fórmula recomendada",
        "TCO mensual = plataforma + comisiones + integrador + multicanal + aplicaciones + desarrollo/soporte + costo operativo manual + impacto esperado de incidentes. A esto se suma la migración inicial amortizada en 24 meses.",
        PALE_BLUE,
        BLUE,
    )
    add_text(
        doc,
        "Con ventas bajas, además del TCO absoluto deben seguirse dos indicadores: costo tecnológico como porcentaje de las ventas netas y cantidad de pedidos necesaria para cubrir el costo fijo digital. Esto permite demostrar si la arquitectura actual está sobredimensionada para el nivel de actividad.",
    )
    for text in (
        "Carga tecnológica = TCO mensual / ventas netas mensuales.",
        "Pedidos de equilibrio tecnológico = TCO mensual / margen de contribución promedio por pedido.",
        "Ahorro de caja = TCO actual - TCO alternativo - amortización mensual de la migración.",
    ):
        add_bullet(doc, text)

    rows = [
        ("VTEX", "Cotización/contrato vigente", "Licencia y condiciones comerciales; no usar estimaciones genéricas"),
        ("Producteca", "Cotización/contrato vigente", "Implementación, customer success/KAM, fee por operación/canales/usuarios y desarrollos según términos"),
        ("Pirce", "Contrato y horas reales", "Abono, alcance, bolsas de horas, incidentes, desarrollos y dependencia"),
        ("Tiendanube Evolución", "Cotización integral", "Fee, transacciones, Pago Nube, apps, desarrollo custom, migración, soporte y condiciones preferenciales"),
        ("Shopify", "Plan USD + pagos externos + apps", "Plan, cargo de Shopify por proveedor externo, gateway local, tema, aplicaciones y desarrollo"),
        ("Astroselling", "Plan + órdenes + posible API", "Canales, SKU, órdenes excedentes, ERP/facturación, API y soporte"),
        ("Migración", "Proyecto único amortizado", "Catálogo, SEO, diseño, contenidos, clientes, promociones, analytics, QA, capacitación y doble operación"),
    ]
    add_table(doc, ["Partida", "Dato requerido", "Qué incluir"], rows, [1800, 2450, 5110], font_size=8.3, first_col_bold=True)

    doc.add_heading("Referencias públicas al 18/08/2026", level=2)
    for text in (
        "Tiendanube Evolución no publica un precio fijo: se contrata mediante cotización. La propuesta debe detallar fee, costo por transacción, condiciones de Pago Nube, apps, desarrollo, migración y soporte. Evolución publica especialista dedicado, atención prioritaria, migración asistida, personalización avanzada y hasta tres centros de distribución.",
        "Shopify Argentina publica, con pago anual: Basic USD 19/mes, Grow USD 49/mes y Advanced USD 299/mes. Con proveedores de pago externos publica cargos adicionales de 2%, 1% y 0,6%, respectivamente. Argentina no figura en la lista consultada de Shopify Payments, por lo que debe presupuestarse un proveedor externo.",
        "Astroselling Argentina publica, con pago anual equivalente: Starter USD 22,42/mes, Pro USD 49,92/mes y Business USD 99,92/mes; también muestra valores mensuales superiores para Pro y Business. Los planes tienen límites de canales/SKU y cargos por órdenes excedentes.",
        "Producteca no muestra un precio simple comparable: sus términos describen implementación, servicios de customer success/KAM, licenciamiento según ventas/canales/usuarios y eventuales horas de desarrollo.",
    ):
        add_bullet(doc, text)
    add_small_source(doc, "Valores orientativos y volátiles. Revalidar con propuestas comerciales antes de decidir.")

    doc.add_heading("Umbral sugerido para justificar la migración", level=2)
    add_text(
        doc,
        "Dado el nivel bajo de ventas, priorizar alternativas que reduzcan al menos 25% del TCO actual y recuperen la inversión de migración dentro de 12 meses, sin dejar requisitos críticos sin solución ni degradar la experiencia visual. Estos porcentajes son criterios de gestión propuestos, no benchmarks del proveedor.",
    )
    add_text(
        doc,
        "Como Evolución es el nivel personalizado de Tiendanube, no debe usarse el precio de Esencial, Impulso o Escala para proyectar ahorro. La comparación válida requiere una cotización Evolución B2C completa y vinculante.",
    )


def section_fit(doc: Document) -> None:
    doc.add_heading("5. Cuándo conviene cada opción", level=1)
    doc.add_heading("Conviene continuar con VTEX cuando...", level=2)
    for text in (
        "Las promociones, el checkout, los inventarios o el fulfillment tienen reglas que Tiendanube Evolución y Shopify no resuelven sin múltiples apps o desarrollo.",
        "Una solución acotada permite resolver económicamente la condición 4x3 + 6 cuotas y simplificar su administración.",
        "La app VTEX IO de cuotas/promociones es crítica y no existe reemplazo equivalente.",
        "La operación proyecta más canales, centros, marcas, B2B, sellers o una experiencia altamente personalizada.",
        "La estabilidad del flujo Producteca-VTEX es buena y el ahorro potencial no compensa el riesgo de migración.",
    ):
        add_bullet(doc, text)

    doc.add_heading("Conviene evaluar Tiendanube Evolución + Astroselling cuando...", level=2)
    for text in (
        "La propuesta Evolución está alineada con el negocio B2C, la tienda propia y Mercado Libre sin pagar por complejidad innecesaria.",
        "La prioridad es reducir costo fijo, dependencia técnica y tiempo administrativo.",
        "El especialista dedicado, soporte y migración asistida compensan la dotación interna reducida.",
        "La dotación reducida puede operar promociones, catálogo y pedidos desde flujos simples y documentados.",
        "La demo resuelve de punta a punta el 4x3 limitado a 6 cuotas o propone una mecánica comercial equivalente aceptable.",
        "Los requerimientos críticos quedan cubiertos por el plan elegido y aplicaciones con soporte/SLA aceptables.",
        "El ahorro a 24 meses supera migración, capacitación, doble operación y recreación de funciones.",
    ):
        add_bullet(doc, text)

    doc.add_heading("Conviene evaluar Shopify + Astroselling cuando...", level=2)
    for text in (
        "El recorrido B2C, diseño, experiencia móvil y ecosistema de aplicaciones superan a las otras alternativas en la prueba.",
        "La calidad visual, la flexibilidad del theme y el ecosistema de aplicaciones pesan más que el acompañamiento y adaptación local de Tiendanube Evolución.",
        "El TCO en USD, el gateway externo y los cargos por transacción siguen siendo menores que la arquitectura actual.",
        "La promoción 4x3 + 6 cuotas puede resolverse y administrarse sin convertir al equipo en dependiente permanente de desarrollo.",
        "Astroselling demuestra catálogo, stock, precio y pedidos con Shopify sobre los SKU reales de Saracho.",
    ):
        add_bullet(doc, text)

    doc.add_heading("Caso hipotético VTEX + Astroselling: cuándo descartarlo", level=2)
    for text in (
        "La conexión VTEX depende solo de una API genérica o trabajo a medida sin SLA de punta a punta.",
        "No se puede demostrar sincronización bidireccional de variantes, atributos, precios, stocks y pedidos.",
        "El ahorro de licencia se traslada a más horas de Pirce, conciliación manual o riesgo de sobreventa.",
    ):
        add_bullet(doc, text)


def section_migration_risks(doc: Document) -> None:
    doc.add_heading("6. Riesgos de migración y controles", level=1)
    rows = [
        ("Identidad de SKU", "Enlaces incorrectos, duplicados o stock cruzado", "Maestro único por SKU/EAN; tabla de equivalencias; prohibir ediciones simultáneas"),
        ("Catálogo de neumáticos", "Pérdida de atributos, filtros, kits x2/x4, dimensiones o SEO", "Exportación completa; muestra de QA; reconciliación por campos y variantes"),
        ("Stock/precio", "Sobreventa, precio desactualizado o margen incorrecto", "Pruebas de concurrencia; alertas; comparación automática; stock de seguridad"),
        ("Pedidos", "Órdenes incompletas, duplicadas o sin facturación", "Idempotencia; estados mapeados; reintentos; conciliación diaria"),
        ("Promociones/cuotas", "Aplicar 4x3 con financiación incorrecta o comunicar condiciones inconsistentes", "Probar 4x3 solo en 6 cuotas; combinaciones, exclusiones, visualización y pedido final"),
        ("SEO y URLs", "Caída de tráfico orgánico", "Mapa 301, metadatos, canonicals, Search Console y monitoreo 90 días"),
        ("Clientes y datos", "Históricos incompletos o problemas de privacidad", "Definir qué migra, consentimiento, seguridad, retención y acceso"),
        ("Responsabilidades", "Incidentes rebotan entre plataforma, app e integrador", "Matriz RACI y SLA de punta a punta con un dueño por dato"),
        ("Dotación reducida", "El equipo no absorbe controles, errores o varios paneles", "Medir horas y pasos; automatizar; manual operativo; soporte con dueño único"),
        ("Cutover", "Interrupción de venta o imposibilidad de volver atrás", "Doble operación limitada, congelamiento, checklist y rollback probado"),
    ]
    add_table(doc, ["Riesgo", "Impacto", "Control mínimo"], rows, [1750, 2900, 4710], font_size=7.8, first_col_bold=True)

    add_callout(
        doc,
        "Riesgo específico de Saracho",
        "El catálogo de neumáticos combina medidas, variantes, marcas, kits, dimensiones logísticas y datos SEO. El identificador de enlace no puede improvisarse: SKU de referencia y EAN deben estar completos y ser estables antes del piloto.",
        PALE_AMBER,
        AMBER,
    )


def section_plan(doc: Document) -> None:
    doc.add_page_break()
    doc.add_heading("7. Plan de decisión recomendado", level=1)
    steps = [
        ("1. Medir la base actual — 2 semanas", "Consolidar facturas, GMV, pedidos, SKU, horas del equipo, tareas manuales, incidentes, latencia de sincronización y alcance real de Pirce."),
        ("2. Cerrar requisitos críticos — 1 semana", "Documentar el recorrido B2C, búsqueda por medida, matriz promocional 4x3 + 6 cuotas, pagos, envíos, Mercado Libre, SEO, reportes y customizaciones."),
        ("3. Solicitar respuestas vinculantes — 1 a 2 semanas", "Obtener una cotización Tiendanube Evolución completa y pedir a todos los proveedores alcance, límites, SLA, seguridad, soporte, costos y responsabilidad por incidentes."),
        ("4. Ejecutar pruebas técnica, promocional y visual — 30 días", "Probar 50 a 100 SKU y dos canales. Construir demos en Tiendanube Evolución y Shopify; ejecutar el recorrido B2C y el 4x3 solo en 6 cuotas hasta el pedido final."),
        ("5. Armar business case — 1 semana", "Comparar TCO a 24 meses, ahorro, repago, riesgos, horas del equipo, brechas y costo de transición."),
        ("6. Decidir y, si corresponde, migrar — 6 a 12 semanas", "Migración por oleadas, QA, redirecciones SEO, capacitación, observabilidad, doble operación acotada y rollback."),
    ]
    for title, detail in steps:
        p = doc.add_paragraph(style="List Number")
        r = p.add_run(title)
        set_font(r, bold=True, color=NAVY)
        r = p.add_run(". " + detail)
        set_font(r)

    doc.add_heading("Criterios propuestos de aceptación del piloto", level=2)
    rows = [
        ("Stock", "Diferencia <= 0,5% de SKU medidos; cero sobreventas atribuibles al integrador"),
        ("Precios", "Diferencia <= 0,5%; reglas de markup verificadas por canal"),
        ("Sincronización", "95% de cambios reflejados en <= 5 minutos o SLA acordado mejor"),
        ("Pedidos", "100% importados una sola vez, con estados y datos necesarios"),
        ("Catálogo", "100% de campos críticos, imágenes, variantes y kits validados en la muestra"),
        ("Operación", "Reducción demostrable de tareas manuales y trazabilidad de errores"),
        ("Promociones", "4x3 se aplica únicamente con 6 cuotas; otras financiaciones quedan excluidas y el pedido conserva la condición"),
        ("Dotación", "Una persona puede crear, pausar, verificar y auditar la promo y la operación rutinaria con un procedimiento claro"),
        ("Soporte", "Escalamiento probado, tiempos de respuesta y dueño de punta a punta definidos"),
        ("Diseño custom", "Demos Tiendanube Evolución y Shopify: Home, PLP, PDP, cuotas, carrito y móvil igualan o mejoran la identidad actual"),
        ("Conversión", "No aparecen fricciones nuevas en búsqueda, selección por medida, compra, pagos o envío"),
        ("Recorrido B2C", "Un consumidor puede encontrar su medida, entender cuotas/promoción y comprar desde móvil sin asistencia"),
    ]
    add_table(doc, ["Métrica", "Condición de aprobación"], rows, [2100, 7260], font_size=8.6, first_col_bold=True)
    add_small_source(doc, "Los umbrales son criterios de aceptación propuestos para Saracho, no compromisos publicados de los proveedores.")

    doc.add_heading("Regla de decisión", level=2)
    add_callout(
        doc,
        "Recomendación final",
        "Para este negocio B2C, comparar VTEX + Producteca, Tiendanube Evolución + Astroselling y Shopify + Astroselling. Elegir la que reduzca materialmente el costo mensual, permita operar el 4x3 solo en 6 cuotas, sea manejable con poca dotación y maximice conversión móvil. Evolución debe evaluarse con cotización real; VTEX + Astroselling queda fuera del caso base.",
        PALE_GREEN,
        GREEN,
    )


def section_vendor_questions(doc: Document) -> None:
    doc.add_heading("8. Preguntas que deben responder los proveedores", level=1)
    doc.add_heading("Astroselling", level=2)
    for text in (
        "Para el caso hipotético VTEX: ¿pueden formalizar la disposición conversada, con alcance, precio, cronograma, mantenimiento y responsables?",
        "Para Tiendanube Evolución y Shopify: ¿qué alcance productivo estándar está soportado y qué operaciones/referencias pueden demostrar?",
        "¿Sincroniza en ambos sentidos catálogo, variantes, atributos, imágenes, precio, promociones, stock y pedidos?",
        "¿La API está incluida o tiene costo de configuración? ¿Quién mantiene el conector ante cambios de VTEX?",
        "¿Qué SLA, monitoreo, reintentos, historial, alertas y soporte ofrece para Argentina?",
        "¿Cómo maneja combos/kits, múltiples depósitos y reglas de markup/stock por canal?",
    ):
        add_bullet(doc, text)

    doc.add_heading("Tiendanube Evolución", level=2)
    for text in (
        "¿Cuál es la cotización Evolución completa para Saracho y qué incluye en fee, transacciones, Pago Nube, migración, soporte, apps y desarrollo?",
        "¿Cómo se optimiza el recorrido B2C móvil: búsqueda por medida, PDP, cuotas, promoción, checkout, pago y envío?",
        "¿Cómo se reemplaza la app de cuotas/promociones construida en VTEX IO?",
        "¿Cómo se implementa un 4x3 que sea válido únicamente con 6 cuotas y se excluya con cualquier otra financiación?",
        "¿Qué límites de API, aplicaciones a medida, exportaciones, redirecciones y migración de clientes aplican?",
        "¿Cuál es el costo efectivo de Evolución con el medio de pago elegido y el GMV real?",
        "¿Qué acceso y límites existen para un desarrollo visual custom que replique o mejore Home, PLP, PDP, cuotas, carrito y experiencia móvil de VTEX?",
    ):
        add_bullet(doc, text)

    doc.add_heading("Shopify", level=2)
    for text in (
        "¿Qué plan, gateway argentino, tema y aplicaciones forman el costo total para Saracho?",
        "¿Cómo se implementa el 4x3 válido únicamente con 6 cuotas: app existente, Shopify Functions o desarrollo a medida?",
        "¿Qué parte del checkout puede personalizarse en el plan propuesto y qué requeriría Shopify Plus?",
        "¿Cómo se integra Astroselling para catálogo, stock, precios y pedidos, y qué soporte cubre cada proveedor?",
        "¿Una persona puede administrar promociones y operación rutinaria sin intervención técnica recurrente?",
        "¿Cómo se optimiza y mide el recorrido B2C móvil para búsqueda por medida, cuotas, pagos y envío en Argentina?",
    ):
        add_bullet(doc, text)

    doc.add_heading("Producteca y Pirce", level=2)
    for text in (
        "¿Qué funciones se usan realmente y cuáles podrían eliminarse o renegociarse?",
        "¿Quién es el dueño operativo de catálogo, stock, precio, pedidos e incidentes?",
        "¿Qué SLA y reportes de errores/incidentes se cumplen hoy?",
        "¿Cuál sería el costo de continuidad optimizada y cuál el de una migración asistida?",
        "¿Qué parte del fee de Pirce corresponde a complejidad de VTEX y qué parte seguiría existiendo en Tiendanube Evolución o Shopify?",
        "¿Puede la solución actual implementar 4x3 solo en 6 cuotas? ¿Con qué costo, plazo y facilidad de administración?",
    ):
        add_bullet(doc, text)


def section_sources(doc: Document) -> None:
    doc.add_page_break()
    doc.add_heading("9. Fuentes y evidencia", level=1)
    add_text(
        doc,
        "Se priorizaron páginas oficiales de los proveedores y documentación interna del proyecto. Consulta web realizada el 18 de agosto de 2026.",
    )
    sources = [
        ("VTEX Developers — Composability", "https://developers.vtex.com/docs/guides/composability"),
        ("VTEX — Comercio B2C omnicanal", "https://www.vtex.com/es-mx/solutions/business-needs/b2c-omnichannel/"),
        ("Producteca — Integración con VTEX", "https://www.producteca.com/integraciones/vtex"),
        ("Producteca — Plataforma e integraciones", "https://www.producteca.com/"),
        ("Producteca — Términos, esquema de precios y SLA", "https://www.producteca.com/terminos-y-condiciones-de-uso-y-politicas-de-privacidad"),
        ("Tiendanube — Planes y precios Argentina", "https://www.tiendanube.com/planes-y-precios"),
        ("Tiendanube Evolución — Plataforma, diseño y migración", "https://www.tiendanube.com/evolucion/plataforma"),
        ("Tiendanube Evolución — Servicios y soporte", "https://www.tiendanube.com/evolucion/servicios"),
        ("Tiendanube Ayuda — Costos y transacciones", "https://ayuda.tiendanube.com/es_AR/sobre-tiendanube/que-costo-tiene-tiendanube"),
        ("Tiendanube Ayuda — Integrar Astroselling", "https://ayuda.tiendanube.com/es_AR/123518-aplicaciones-de-canales-de-venta/como-integrar-astroselling-con-mi-tienda"),
        ("Astroselling Argentina — Funciones, integraciones y planes", "https://www.astroselling.com/es/ar/"),
        ("Astroselling Ayuda — Conectores de e-commerce", "https://help.astroselling.com/es/collections/3653305-conectar-e-commerce"),
        ("Shopify Argentina — Planes, cargos por pagos externos y funciones", "https://www.shopify.com/ar/precios"),
        ("Shopify Ayuda — Países admitidos para Shopify Payments", "https://help.shopify.com/es/manual/payments/shopify-payments/supported-countries"),
        ("Shopify Ayuda — Descuentos Compra X y obtén Y", "https://help.shopify.com/es/manual/discounts/discount-types/buy-x-get-y"),
        ("Shopify Developers — Extensión de descuentos con Shopify Functions", "https://shopify.dev/docs/apps/build/discounts/index"),
    ]
    for title, url in sources:
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.space_after = Pt(4)
        add_hyperlink(p, title, url)

    doc.add_heading("Evidencia interna consultada", level=2)
    for text in (
        "REGLAS_PROYECTO.md: confirma que la tienda usa VTEX y que Producteca sincroniza el catálogo con VTEX.",
        "Repositorio installments-info-sarachoneumaticos: evidencia una aplicación personalizada VTEX IO para información de cuotas/promociones.",
        "Archivos y reportes del proyecto: evidencian catálogo de neumáticos, kits, dimensiones logísticas y reportes operativos VTEX.",
        "Antecedente informado por Saracho Neumáticos: ya hubo conversaciones con Astroselling y el proveedor manifestó disposición a desarrollar una integración con VTEX.",
        "Objetivo informado por Saracho Neumáticos: reducir costos por el nivel bajo de ventas y demostrar mediante desarrollo custom que Tiendanube Evolución o Shopify pueden alcanzar una presentación comparable con VTEX.",
        "Restricción informada por Saracho Neumáticos: existen promociones como 4x3 válidas solo en 6 cuotas que la implementación actual de VTEX no permite limitar como se necesita.",
        "Capacidad informada por Saracho Neumáticos: la dotación dedicada a la página web es reducida y no puede absorber alta carga manual ni coordinación constante.",
        "Modelo informado por Saracho Neumáticos: el negocio es B2C; deben priorizarse conversión, móvil, cuotas, promociones, pagos, envíos y autogestión.",
        "Alternativa informada por Saracho Neumáticos: la propuesta Tiendanube a evaluar corresponde específicamente al plan Evolución.",
    ):
        add_bullet(doc, text)

    doc.add_heading("Lectura de la integración VTEX propuesta por Astroselling", level=2)
    add_text(
        doc,
        "La documentación pública consultada de Astroselling enumera conectores como Tiendanube, Shopify, WooCommerce, Wix y otros, pero no VTEX. Aunque Astroselling manifestó disposición a desarrollar esa integración, el escenario VTEX + Astroselling se mantiene como hipótesis y queda fuera del caso base hasta que exista propuesta formal y prueba de concepto.",
    )

    add_callout(
        doc,
        "Nota final",
        "Este informe recomienda un proceso de decisión; no una migración inmediata. La información comercial y técnica debe cerrarse con propuestas vigentes, contratos, SLAs y pruebas sobre el catálogo real de Saracho Neumáticos.",
        LIGHT,
        NAVY,
    )


def build() -> Path:
    doc = Document()
    set_doc_defaults(doc)
    for section in doc.sections:
        set_running_header_footer(section)
    title_page(doc)
    section_context(doc)
    section_platform_comparison(doc)
    section_ecosystems(doc)
    section_costs(doc)
    section_fit(doc)
    section_migration_risks(doc)
    section_plan(doc)
    section_vendor_questions(doc)
    section_sources(doc)

    # Document metadata and pagination behavior.
    doc.core_properties.title = "VTEX vs. Tiendanube vs. Shopify — Saracho Neumáticos"
    doc.core_properties.subject = "Evaluación de plataforma e integradores"
    doc.core_properties.author = "Saracho Neumáticos"
    doc.core_properties.keywords = "VTEX, Tiendanube, Shopify, Producteca, Astroselling, Pirce, ecommerce"

    for paragraph in doc.paragraphs:
        if paragraph.style and paragraph.style.name.startswith("Heading"):
            set_keep_with_next(paragraph, True)
        if len(paragraph.text) < 180:
            set_keep_together(paragraph, True)

    doc.save(OUTPUT)
    return OUTPUT


if __name__ == "__main__":
    print(build())
